from pathlib import Path
from decimal import Decimal as D
import hashlib, json, re, zipfile
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / 'research/raw'
KIT = ROOT / 'output/starter-kit/sources'
OUT = ROOT / 'output/comparability-audit'
OUT.mkdir(parents=True, exist_ok=True)

rows = [
 (1,'Brazil','Kinross / Paracatu','parked',None,'Current registry and historical SEC report do not establish the same construction stage. The last-raise-date bridge fails a cross-check against SEC-reported Eustaquio Stage 15 completion.', ['kinross-paracatu-2026.html','brazil-dams.csv','audit-sigbm-manual.pdf']),
 (2,'Peru','Southern Copper / Toquepala + Cuajone','retained','Calendar 2023; fine molybdenum output; full scope of both Peruvian mines','Retained historical 2023 from the 2024 annual workbook. Its 2024 column is explicitly preliminary and excluded.', ['southern-sec-2024.html','peru-dec2024.xlsx','audit-peru-annual-2024.xlsx','audit-southern-memoria2024.pdf']),
 (3,"Cote d'Ivoire",'Endeavour / Ity','retained','Calendar 2024; SMI+SMD+SMF; both monetary mining royalty components; 100% project payments','AIF explicitly maps all three licence-holding companies. Company-initial and final royalty cells agree; comparison uses the disclosed FX range.', ['endeavour-estma-2024.pdf','ivory-annexes.bin','ivory-2024.pdf','audit-endeavour-aif-2024.pdf']),
 (4,'Senegal','Endeavour / Sabodala-Massawa','parked',None,'ESTMA project label Sabodala Mining Company SARL is not decisively bridged to the host declaration by separate legal entity Sabodala Gold Operations. 2025 government-side royalties were already removed in reconciliation.', ['endeavour-estma-2024.pdf','senegal-extra.bin','senegal-2024.pdf','audit-endeavour-aif-2024.pdf']),
 (5,'Mali','B2Gold / Fekola','parked',None,'Gross/net dividends and ordinary-shareholder versus qualifying distributions remain unresolved. Royalty basket mapping to ad-valorem plus ISCP is not explicit enough.', ['b2gold-estma-2023.pdf','mali-annexes-2023.xlsx','mali-2023.pdf']),
 (6,'Madagascar','Rio Tinto / QMM','parked',None,'UK annual-average FX basis is explicit but exact rate, underlying currencies, and royalty/recipient basket are not aligned with EITI.', ['uk-rio-2023.json','rio-fca-payments-2023.html','madagascar-2023.pdf']),
 (7,'DRC','Barrick / Kibali','retained','2023; Barrick group economic interest of 45%, not immediate registered shareholding','2023 ownership form and affirmative company-supply marker; host explicitly attributes Barrick 45% to company declaration; Western AIF applies 45% to 2023 operations.', ['barrick-sec-aif-2023.html','congo-annexes.bin','congo-2023.pdf']),
 (8,'Burkina Faso','IAMGOLD / Essakane','retained','Group/state economic interests 90%/10% across admissible 2023-2024 date interpretations','Original embedded form has conflicting date labels. Both interpretations precede the documented 20 June 2025 ownership change; no date label is silently corrected.', ['iamgold-sec-aif-2023.html','iamgold-sec-2024.html','audit-iamgold-sec-2025.html','burkina-annexes.bin','burkina-2024.pdf']),
 (9,'Zambia','Barrick / Lumwana','parked',None,'Reconciled royalty differs from original company figure; no unambiguous final company adoption or matching net-cash basis. Broad Western taxes are not automatically the host income-tax field.', ['barrick-estma-2023.pdf','zambia-2023-2024.pdf']),
 (10,'Ghana','Galiano / Asanko','parked',None,'Host GRA/MIIF mineral-royalty basket is not proved to include the separate Western Register General recipient. Currency notes do not cover the whole Western total.', ['galiano-estma-2023.pdf','ghana-2023.pdf']),
]
cases = [dict(id=f'RAB-MIN-{n:02}',country=c,project=p,decision=d,approved_scope=s,reason=r,source_ids=ids,proven_contradiction=False) for n,c,p,d,s,r,ids in rows]

# Recompute quantitative controls from the inspected originals.
sec = re.sub(r'\s+', ' ', (RAW/'southern-sec-2024.txt').read_text())
segments = [sec[m.start():m.start()+1600] for m in re.finditer(re.escape('MOLYBDENUM (thousand pounds)'), sec)]
segment = next(s for s in segments if 'Toquepala' in s[:500] and 'Cuajone' in s[:500])
mines = {}
for mine in ['Toquepala','Cuajone']:
    m = re.search(mine + r'\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)', segment)
    if not m: raise ValueError(f'Molybdenum table did not parse: {mine}')
    mines[mine] = dict(zip(['2024','2023','2022'],[int(x.replace(',','')) for x in m.groups()]))
peru = openpyxl.load_workbook(RAW/'audit-peru-annual-2024.xlsx',read_only=True,data_only=True)['Mo-Empresas']
converted = sum(D(mines[m]['2023']) for m in mines)*D('0.45359237')
halfwidth = D('0.90718474')  # one 1,000-lb reporting unit per mine, conservative for rounding/truncation
host = D(str(peru['J7'].value))

ivory = openpyxl.load_workbook(KIT/'ivory-annexes.xlsx',read_only=True,data_only=True)
cells = []
for sn in ['34.1','34.5','34.8']:
    s=ivory[sn]
    for row in [27,42]:
        cells.append(dict(sheet=sn,company=s['E2'].value,initial_cell=f'E{row}',final_cell=f'G{row}',label=s[f'C{row}'].value,initial=s[f'E{row}'].value,adjustment=s[f'F{row}'].value,final=s[f'G{row}'].value))
ity = sum(D(str(c['final'])) for c in cells)
ity_low,ity_high=ity/D('631.40'),ity/D('595.72')

congo = openpyxl.load_workbook(KIT/'congo-annexes.xlsx',read_only=True,data_only=True)
burkina = openpyxl.load_workbook(RAW/'burkina-embedded-Microsoft_Excel_Worksheet3.xlsx',read_only=True,data_only=True)['17.Structure de Capital ']
controls = {
 'peru_2023':dict(sec_thousand_lb=mines,host_cell='Mo-Empresas!J7',host_tonnes=str(host),normalized_tonnes=str(converted),rounding_interval_tonnes=[str(converted-halfwidth),str(converted+halfwidth)],consistent=converted-halfwidth<=host<=converted+halfwidth,excluded_2024_header=peru['K4'].value,preliminary_note=peru['A16'].value),
 'ity_2024':dict(cells=cells,host_xof=str(ity),fx_xof_per_usd=['595.72','631.40'],usd_interval_before_precision_allowance=[str(ity_low),str(ity_high)],western_usd='45145203',consistent=ity_low<=D('45145203')<=ity_high),
 'kibali_2023':dict(host_reporting_period=congo['Annexe 10']['A6'].value,company_submitted=congo['Annexe 11 ']['C29'].value,identifier=congo['Annexe 11 ']['F29'].value,company_attribution=congo['Annexe 13 ']['D27'].value,host_barrick_percent=congo['Annexe 11 ']['N29'].value,western_barrick_percent=45,consistent=congo['Annexe 11 ']['N29'].value==45),
 'essakane_date_bracket':dict(period_header=burkina['A6'].value,capital_date_label=burkina['A12'].value,company=burkina['C15'].value,group_fraction=burkina['D15'].value,state_fraction=burkina['D13'].value,western_prior_to_change=[0.9,0.1],change_effective='2025-06-20',consistent=burkina['D15'].value==0.9 and burkina['D13'].value==0.1),
 'scope_note':'These calculations verify extraction and arithmetic. The legal/entity/period crosswalks are documented research judgments, not proven by numerical agreement.'
}
assert all(controls[k]['consistent'] for k in ['peru_2023','ity_2024','kibali_2023','essakane_date_bracket'])

logs=[json.loads(x) for x in (ROOT/'research/access-log.jsonl').read_text().splitlines()]
by_id={x.get('id'):x for x in logs}
ids=sorted(set(x for c in cases for x in c['source_ids']))
manifest=[]
for sid in ids:
    p=RAW/sid
    entry=dict(source_id=sid,**{k:by_id.get(sid,{}).get(k) for k in ['url','final_url','tested_at_utc','status','content_type']})
    if p.exists():
        data=p.read_bytes();entry.update(bytes=len(data),sha256=hashlib.sha256(data).hexdigest())
    else: entry['local_bytes_found']=False
    manifest.append(entry)

result=dict(audit_date='2026-10-12',retained=4,parked=6,proven_contradictions=0,cases=cases)
(OUT/'audit-results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
(OUT/'numeric-controls.json').write_text(json.dumps(controls,ensure_ascii=False,indent=2)+'\n')
(OUT/'source-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
(OUT/'audit-report.md').write_text((ROOT/'research/audit-report.md').read_text())
(OUT/'README.md').write_text('Rabbit mining comparability audit, 12 October 2026\n\nFour retained narrow comparisons; six parked; no proven contradictions. Read audit-report.md first. audit-results.json records every verdict; numeric-controls.json preserves the checked values and calculations; source-manifest.json gives retrieval URLs and hashes for inspected originals. This is a decision/evidence-index packet, not an archive of all source documents. A retained comparison does not authorize comparison of other fields, dates or company scopes.\n')
archive=ROOT/'output/Rabbit-mining-comparability-audit.zip'
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
    for p in sorted(OUT.iterdir()):z.write(p,p.name)
print(json.dumps(dict(retained=4,parked=6,controls_consistent=True,source_records=len(manifest),sources_without_url=[x['source_id'] for x in manifest if not x.get('url')],missing_local_sources=[x['source_id'] for x in manifest if x.get('local_bytes_found') is False],archive=str(archive),bytes=archive.stat().st_size)))
